from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.colors import HexColor
from reportlab.pdfgen.canvas import Canvas


ROOT = Path(__file__).resolve().parents[1]


def set_font(run, name="Calibri", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_font(r, size=16 if level == 1 else 13, bold=True, color="2E74B5")
    return p


def build_docx():
    out = ROOT / "submissions" / "CSI473-Lab-01-Hand-in.docx"
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("CSI473 Laboratory 1 Hand-in")
    set_font(run, size=23, bold=True, color="0B2545")
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(16)
    set_font(sub.add_run("Campus Concierge - Studio setup and project launch"), size=13, color="555555")

    for label, value in [("Repository", "CampusConcierge-"), ("Evidence", "Repository link, editable model, glossary, candidate problem and design record"), ("Status", "Draft - complete reviewer details after peer check")]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        set_font(p.add_run(label + ": "), bold=True)
        set_font(p.add_run(value))

    add_heading(doc, "Candidate project problem")
    doc.add_paragraph("University of Botswana students often need help from several administrative and support units, but information is spread across different channels and it is difficult to know who owns an enquiry or whether it has progressed. Campus Concierge will provide one accessible entry point where students can search approved guidance, submit a structured request, receive a clear route to the responsible service owner, and follow its status. The project must protect personal information, keep service ownership visible, and remain usable on low-bandwidth mobile connections.")

    add_heading(doc, "Exit record")
    items = [
        ("Evidence that changed our mind", "The brief requires editable, reproducible source evidence rather than screenshots."),
        ("Decision made", "Organise the repository around design evidence and keep an editable SVG context model with exported SVG/PDF versions."),
        ("Consequence accepted", "The team must maintain source and export files and record the evidence behind design decisions."),
        ("First risk", "Service owners may disagree about request categories and routing responsibilities."),
        ("Next action", "Consult Student Affairs and an academic department to validate service categories and escalation rules."),
    ]
    for label, value in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        set_font(p.add_run(label + ": "), bold=True)
        set_font(p.add_run(value))

    add_heading(doc, "Repository evidence")
    p = doc.add_paragraph()
    p.add_run("Included: ").bold = True
    p.add_run("contribution agreement; glossary template; editable context diagram and PDF export; folders for models, prototype, tests and evidence; reproducibility checklist; and the first design record.")
    doc.save(out)


def box(c, x, y, w, h, fill, stroke, title, subtitle):
    c.setFillColor(HexColor("#" + fill)); c.setStrokeColor(HexColor("#" + stroke)); c.setLineWidth(2)
    c.roundRect(x, y, w, h, 14, fill=1, stroke=1)
    c.setFillColor(HexColor("#1F2937")); c.setFont("Helvetica-Bold", 17)
    c.drawCentredString(x + w / 2, y + h / 2 + 9, title)
    c.setFont("Helvetica", 11)
    c.drawCentredString(x + w / 2, y + h / 2 - 12, subtitle)


def build_pdf():
    out = ROOT / "models" / "context-diagram.pdf"
    c = Canvas(str(out), pagesize=(864, 518))
    box(c, 50, 210, 170, 85, "E8EEF5", "1F4D78", "Students", "search, submit, track")
    box(c, 300, 170, 264, 160, "D9EAD3", "38761D", "Campus Concierge", "guidance, request capture, triage, status")
    box(c, 644, 210, 170, 85, "E8EEF5", "1F4D78", "Service owners", "resolve and update")
    box(c, 340, 50, 184, 65, "FFF2CC", "BF9000", "Knowledge articles", "approved guidance")
    c.setStrokeColor(HexColor("#4B5563")); c.setLineWidth(1.8)
    for x1, y1, x2, y2, label in [(220,252,300,252,"enquiry / request"),(564,252,644,252,"routed request"),(644,228,564,228,"status / resolution"),(432,115,432,170,"approved guidance")]:
        c.line(x1,y1,x2,y2)
        c.setFillColor(HexColor("#4B5563")); c.setFont("Helvetica", 10)
        c.drawCentredString((x1+x2)/2, (y1+y2)/2+12, label)
    c.save()


if __name__ == "__main__":
    build_docx()
    build_pdf()
